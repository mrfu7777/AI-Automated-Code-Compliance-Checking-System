from __future__ import annotations

import re
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any, Literal

import ifcopenshell
import ifcopenshell.util.element
from docx import Document
from openpyxl import load_workbook

from app.services.pdf_extraction import PdfiumRegulationExtractor
from app.services.storage import ObjectStorage

EXTRACTOR_VERSION = "m4.project-extraction.v1"
DocumentKind = Literal["pdf", "docx", "xlsx", "ifc"]


@dataclass(frozen=True)
class ExtractedSource:
    text: str
    location: dict[str, Any]
    confidence: float


@dataclass(frozen=True)
class FactCandidate:
    key: str
    value: Any
    unit: str | None
    scope_data: dict[str, Any]
    location: dict[str, Any]
    excerpt: str
    confidence: float


@dataclass(frozen=True)
class FactDefinition:
    key: str
    label: str
    unit: str | None
    patterns: tuple[str, ...]
    value_kind: Literal["number", "integer", "text", "boolean"] = "number"


FACT_DEFINITIONS = (
    FactDefinition(
        "building.height_m",
        "Building height",
        "m",
        (r"(?:建筑高度|building height)\D{0,12}(\d+(?:\.\d+)?)\s*(?:m|米)",),
    ),
    FactDefinition(
        "building.floor_count",
        "Above-ground floor count",
        "count",
        (r"(?:地上层数|above[- ]ground floors?)\D{0,12}(\d+)",),
        "integer",
    ),
    FactDefinition(
        "building.use",
        "Building use",
        None,
        (r"(?:建筑用途|building use)\s*[:\uff1a=]\s*([^,\uff0c;\uff1b\n]+)",),
        "text",
    ),
    FactDefinition(
        "fire_resistance.rating",
        "Fire resistance rating",
        None,
        (r"(?:耐火等级|fire resistance rating)\s*[:\uff1a=]?\s*([IVX\u2160-\u2163一二三四]+)",),
        "text",
    ),
    FactDefinition(
        "fire_compartment.area_m2",
        "Fire compartment area",
        "m2",
        (r"(?:防火分区面积|fire compartment area)\D{0,12}(\d+(?:\.\d+)?)\s*(?:m2|m²|㎡|平方米)",),
    ),
    FactDefinition(
        "exit.count",
        "Safety exit count",
        "count",
        (r"(?:安全出口数量|number of (?:safety )?exits?)\D{0,12}(\d+)",),
        "integer",
    ),
    FactDefinition(
        "egress.door_clear_width_m",
        "Egress door clear width",
        "m",
        (
            r"(?:疏散门净宽|egress door clear width)\D{0,12}"
            r"(\d+(?:\.\d+)?)\s*(mm|cm|m|毫米|厘米|米)",
        ),
    ),
    FactDefinition(
        "egress.corridor_clear_width_m",
        "Corridor clear width",
        "m",
        (r"(?:疏散走道净宽|corridor clear width)\D{0,12}(\d+(?:\.\d+)?)\s*(mm|cm|m|毫米|厘米|米)",),
    ),
    FactDefinition(
        "egress.stair_clear_width_m",
        "Stair clear width",
        "m",
        (r"(?:疏散楼梯净宽|stair clear width)\D{0,12}(\d+(?:\.\d+)?)\s*(mm|cm|m|毫米|厘米|米)",),
    ),
    FactDefinition(
        "egress.travel_distance_m",
        "Evacuation travel distance",
        "m",
        (r"(?:疏散距离|evacuation travel distance)\D{0,12}(\d+(?:\.\d+)?)\s*(?:m|米)",),
    ),
    FactDefinition(
        "fire_elevator.provided",
        "Fire elevator provided",
        None,
        (
            r"(?:消防电梯|fire elevator)\s*[:\uff1a=]?\s*"
            r"(设置|有|是|provided|yes|true|未设置|无|否|no|false)",
        ),
        "boolean",
    ),
    FactDefinition(
        "sprinkler.provided",
        "Automatic sprinkler provided",
        None,
        (
            r"(?:自动喷水灭火系统|automatic sprinkler)\s*[:\uff1a=]?\s*"
            r"(设置|有|是|provided|yes|true|未设置|无|否|no|false)",
        ),
        "boolean",
    ),
    FactDefinition(
        "smoke_control.provided",
        "Smoke control provided",
        None,
        (
            r"(?:防排烟系统|smoke control)\s*[:\uff1a=]?\s*"
            r"(设置|有|是|provided|yes|true|未设置|无|否|no|false)",
        ),
        "boolean",
    ),
    FactDefinition(
        "fire_alarm.provided",
        "Fire alarm provided",
        None,
        (
            r"(?:火灾自动报警系统|fire alarm)\s*[:\uff1a=]?\s*"
            r"(设置|有|是|provided|yes|true|未设置|无|否|no|false)",
        ),
        "boolean",
    ),
    FactDefinition(
        "hydrant.provided",
        "Fire hydrant provided",
        None,
        (
            r"(?:消火栓系统|fire hydrant)\s*[:\uff1a=]?\s*"
            r"(设置|有|是|provided|yes|true|未设置|无|否|no|false)",
        ),
        "boolean",
    ),
)

_UNIT_ALIASES = {"毫米": "mm", "厘米": "cm", "米": "m"}
_FALSE_VALUES = {"未设置", "无", "否", "no", "false"}


def classify_document(filename: str) -> DocumentKind:
    extension = Path(filename).suffix.lower().lstrip(".")
    if extension not in {"pdf", "docx", "xlsx", "ifc"}:
        raise ValueError(f"Unsupported project document: {extension}")
    return extension  # type: ignore[return-value]


def _candidate_value(definition: FactDefinition, match: re.Match[str]) -> tuple[Any, str | None]:
    raw = match.group(1).strip()
    unit = definition.unit
    if definition.value_kind == "integer":
        return int(raw), unit
    if definition.value_kind == "number":
        if match.lastindex and match.lastindex >= 2:
            unit = _UNIT_ALIASES.get(match.group(2).lower(), match.group(2).lower())
        return float(raw), unit
    if definition.value_kind == "boolean":
        return raw.lower() not in _FALSE_VALUES, None
    return raw, None


def extract_pattern_facts(sources: Iterable[ExtractedSource]) -> list[FactCandidate]:
    candidates: list[FactCandidate] = []
    seen: set[tuple[str, str, str]] = set()
    for source in sources:
        for definition in FACT_DEFINITIONS:
            for pattern in definition.patterns:
                for match in re.finditer(pattern, source.text, flags=re.IGNORECASE):
                    value, unit = _candidate_value(definition, match)
                    identity = (definition.key, repr(value), repr(source.location))
                    if identity in seen:
                        continue
                    seen.add(identity)
                    candidates.append(
                        FactCandidate(
                            key=definition.key,
                            value=value,
                            unit=unit,
                            scope_data={},
                            location=source.location,
                            excerpt=match.group(0).strip(),
                            confidence=source.confidence,
                        )
                    )
    return candidates


class ProjectDocumentExtractor:
    def __init__(self, storage: ObjectStorage) -> None:
        self.storage = storage

    def extract(self, object_key: str, filename: str) -> tuple[DocumentKind, list[FactCandidate]]:
        kind = classify_document(filename)
        if kind == "pdf":
            sources = self._pdf_sources(object_key)
            return kind, extract_pattern_facts(sources)
        with TemporaryDirectory(prefix="code-compliance-m4-") as directory:
            source = Path(directory) / f"source.{kind}"
            self.storage.download_to_file(object_key, source)
            if kind == "docx":
                return kind, extract_pattern_facts(self._docx_sources(source))
            if kind == "xlsx":
                return kind, extract_pattern_facts(self._xlsx_sources(source))
            return kind, self._ifc_candidates(source)

    def _pdf_sources(self, object_key: str) -> list[ExtractedSource]:
        pages = PdfiumRegulationExtractor(self.storage).extract(object_key)
        return [
            ExtractedSource(
                line.text,
                {"page": page.page_number, "bbox": line.bbox, "method": page.method},
                line.confidence,
            )
            for page in pages
            for line in page.lines
        ]

    @staticmethod
    def _docx_sources(path: Path) -> list[ExtractedSource]:
        document = Document(str(path))
        sources = [
            ExtractedSource(paragraph.text, {"paragraph": index + 1}, 1.0)
            for index, paragraph in enumerate(document.paragraphs)
            if paragraph.text.strip()
        ]
        for table_index, table in enumerate(document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if text:
                    sources.append(
                        ExtractedSource(text, {"table": table_index, "row": row_index}, 1.0)
                    )
        return sources

    @staticmethod
    def _xlsx_sources(path: Path) -> list[ExtractedSource]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sources: list[ExtractedSource] = []
        try:
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    populated = [cell for cell in row if cell.value not in {None, ""}]
                    if not populated:
                        continue
                    sources.append(
                        ExtractedSource(
                            " ".join(str(cell.value) for cell in populated),
                            {
                                "sheet": worksheet.title,
                                "range": f"{populated[0].coordinate}:{populated[-1].coordinate}",
                            },
                            1.0,
                        )
                    )
        finally:
            workbook.close()
        return sources

    @staticmethod
    def _ifc_candidates(path: Path) -> list[FactCandidate]:
        model = ifcopenshell.open(path)
        candidates = [
            FactCandidate(
                key=key,
                value=len(model.by_type(ifc_type)),
                unit="count",
                scope_data={},
                location={"ifc_type": ifc_type},
                excerpt=f"{ifc_type} object count",
                confidence=1.0,
            )
            for key, ifc_type in (
                ("building.floor_count", "IfcBuildingStorey"),
                ("space.count", "IfcSpace"),
                ("door.count", "IfcDoor"),
                ("stair.count", "IfcStair"),
            )
        ]
        property_sources: list[ExtractedSource] = []
        for element in model.by_type("IfcObject"):
            global_id = getattr(element, "GlobalId", None)
            try:
                property_sets = ifcopenshell.util.element.get_psets(element)
            except (RuntimeError, TypeError):
                continue
            for set_name, properties in property_sets.items():
                if not isinstance(properties, dict):
                    continue
                for property_name, value in properties.items():
                    if property_name == "id" or isinstance(value, dict | list):
                        continue
                    property_sources.append(
                        ExtractedSource(
                            f"{property_name}: {value}",
                            {
                                "ifc_global_id": global_id,
                                "property_path": f"{set_name}.{property_name}",
                            },
                            1.0,
                        )
                    )
        candidates.extend(extract_pattern_facts(property_sources))
        return candidates
