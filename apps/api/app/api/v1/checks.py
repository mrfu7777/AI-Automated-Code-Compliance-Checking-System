from __future__ import annotations

import hashlib
import io
import json
from typing import Annotated, Any
from uuid import UUID

from fastapi import APIRouter, Depends, Request, status
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import JSONResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen.canvas import Canvas
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.dependencies import Actor, get_current_actor, get_request_id
from app.core.errors import ApplicationError
from app.db.models import (
    CheckResult,
    CheckRun,
    Clause,
    DocumentPage,
    Evidence,
    Job,
    Project,
    ProjectFact,
    ReviewPackage,
    ReviewPackageRulePack,
    Rule,
    RulePack,
    Standard,
    StandardVersion,
    User,
)
from app.db.session import get_database_session
from app.domain.enums import JobStatus
from app.domain.m1_schemas import JobResponse
from app.domain.m3_schemas import (
    CheckResultResponse,
    CheckRunCreate,
    CheckRunCreated,
    CheckRunResponse,
)
from app.domain.m5_schemas import (
    FindingUpdate,
    WorkbenchEvidence,
    WorkbenchFinding,
    WorkbenchResponse,
)
from app.services.audit import record_audit_event
from app.services.dispatch import JobDispatcher, dispatch_persisted_job, get_job_dispatcher
from app.services.rule_engine import ENGINE_VERSION
from app.services.storage import ObjectStorage, get_object_storage

router = APIRouter(tags=["checks"])


def _canonical_hash(value: Any) -> str:
    encoded = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode()).hexdigest()


async def _owned_run(session: AsyncSession, run_id: UUID, actor: Actor) -> CheckRun:
    run = await session.scalar(
        select(CheckRun)
        .join(ReviewPackage)
        .join(Project)
        .where(CheckRun.id == run_id, Project.organization_id == actor.organization_id)
    )
    if run is None:
        raise ApplicationError("check_run_not_found", "Check run was not found", status_code=404)
    return run


async def _run_response(session: AsyncSession, run: CheckRun) -> CheckRunResponse:
    results = list(
        await session.scalars(
            select(CheckResult)
            .where(CheckResult.check_run_id == run.id)
            .order_by(CheckResult.created_at)
        )
    )
    response = CheckRunResponse.model_validate(run)
    return response.model_copy(
        update={"results": [CheckResultResponse.model_validate(item) for item in results]}
    )


async def _workbench_evidence(
    session: AsyncSession, storage: ObjectStorage, evidence_id: str
) -> WorkbenchEvidence | None:
    try:
        evidence_uuid = UUID(evidence_id)
    except ValueError:
        return None
    evidence = await session.get(Evidence, evidence_uuid)
    if evidence is None:
        return None
    image_url = None
    page_number = evidence.location.get("page")
    if evidence.file_version_id and isinstance(page_number, int):
        page = await session.scalar(
            select(DocumentPage).where(
                DocumentPage.file_version_id == evidence.file_version_id,
                DocumentPage.page_number == page_number,
            )
        )
        if page and page.image_object_key:
            image_url = await run_in_threadpool(
                storage.presigned_download,
                page.image_object_key,
                f"page-{page.page_number}.png",
            )
    return WorkbenchEvidence(
        id=evidence.id,
        kind=evidence.kind,
        file_version_id=evidence.file_version_id,
        location=evidence.location,
        excerpt=evidence.excerpt,
        image_url=image_url,
    )


async def _workbench_response(
    session: AsyncSession, storage: ObjectStorage, run: CheckRun
) -> WorkbenchResponse:
    results = list(
        await session.scalars(
            select(CheckResult)
            .where(CheckResult.check_run_id == run.id)
            .order_by(CheckResult.created_at)
        )
    )
    findings = []
    for result in results:
        project_items = [
            item
            for evidence_id in result.project_evidence_ids
            if (item := await _workbench_evidence(session, storage, evidence_id)) is not None
        ]
        regulation_items = [
            item
            for evidence_id in result.regulation_evidence_ids
            if (item := await _workbench_evidence(session, storage, evidence_id)) is not None
        ]
        findings.append(
            WorkbenchFinding(
                result_id=result.id,
                rule_id=result.rule_id,
                status=result.status,
                severity=result.severity,
                message=result.message,
                workflow_status=result.workflow_status,
                assignee_id=result.assignee_id,
                reviewer_notes=result.reviewer_notes,
                trace=result.trace,
                project_evidence=project_items,
                regulation_evidence=regulation_items,
            )
        )
    return WorkbenchResponse(run_id=run.id, findings=findings)


async def _project_snapshot(session: AsyncSession, project_id: UUID) -> dict[str, Any]:
    rows = list(
        await session.scalars(
            select(ProjectFact)
            .where(
                ProjectFact.project_id == project_id,
                ProjectFact.verification_status == "verified",
            )
            .order_by(ProjectFact.created_at.desc())
        )
    )
    superseded_ids = {fact.supersedes_id for fact in rows if fact.supersedes_id is not None}
    latest: dict[tuple[str, str], ProjectFact] = {}
    for fact in rows:
        if fact.id in superseded_ids:
            continue
        scope = json.dumps(fact.scope_data, sort_keys=True, separators=(",", ":"))
        latest.setdefault((fact.key, scope), fact)
    facts = []
    for fact in latest.values():
        evidence_ids = [
            str(item)
            for item in await session.scalars(
                select(Evidence.id).where(Evidence.project_fact_id == fact.id)
            )
        ]
        facts.append(
            {
                "id": str(fact.id),
                "key": fact.key,
                "value": fact.value,
                "unit": fact.unit,
                "scope_data": fact.scope_data,
                "source": fact.source,
                "verified_at": fact.verified_at.isoformat() if fact.verified_at else None,
                "evidence_ids": evidence_ids,
            }
        )
    return {"project_id": str(project_id), "facts": sorted(facts, key=lambda item: item["key"])}


async def _pack_snapshot(session: AsyncSession, packs: list[RulePack]) -> list[dict[str, Any]]:
    result = []
    for pack in sorted(packs, key=lambda item: str(item.id)):
        rules = list(
            await session.scalars(
                select(Rule).where(Rule.rule_pack_id == pack.id).order_by(Rule.code)
            )
        )
        snapshots = []
        for rule in rules:
            clause = await session.get_one(Clause, rule.source_clause_id)
            evidence_ids = [
                str(item)
                for item in await session.scalars(
                    select(Evidence.id).where(Evidence.clause_id == clause.id)
                )
            ]
            snapshots.append(
                {
                    "id": str(rule.id),
                    "code": rule.code,
                    "title": rule.title,
                    "severity": rule.severity,
                    "applicability": rule.applicability,
                    "inputs": rule.inputs,
                    "expression": rule.expression,
                    "missing_data_status": rule.missing_data_status,
                    "clause": {
                        "id": str(clause.id),
                        "number": clause.clause_number,
                        "original_text": clause.original_text,
                        "page_number": clause.page_number,
                        "evidence_ids": evidence_ids,
                    },
                }
            )
        result.append(
            {
                "id": str(pack.id),
                "name": pack.name,
                "semantic_version": pack.semantic_version,
                "content_hash": pack.content_hash,
                "rules": snapshots,
            }
        )
    return result


@router.post("/check-runs", response_model=CheckRunCreated, status_code=status.HTTP_201_CREATED)
async def create_check_run(
    payload: CheckRunCreate,
    request: Request,
    actor: Annotated[Actor, Depends(get_current_actor)],
    session: Annotated[AsyncSession, Depends(get_database_session)],
    dispatcher: Annotated[JobDispatcher, Depends(get_job_dispatcher)],
) -> CheckRunCreated:
    project = await session.scalar(
        select(Project).where(
            Project.id == payload.project_id, Project.organization_id == actor.organization_id
        )
    )
    if project is None:
        raise ApplicationError("project_not_found", "Project was not found", status_code=404)
    pack_ids = list(dict.fromkeys(payload.rule_pack_ids))
    packs = list(
        await session.scalars(
            select(RulePack)
            .join(StandardVersion)
            .join(Standard)
            .where(
                RulePack.id.in_(pack_ids),
                RulePack.lifecycle_status == "published",
                Standard.organization_id == actor.organization_id,
            )
        )
    )
    if len(packs) != len(pack_ids):
        raise ApplicationError(
            "published_rule_pack_not_found",
            "Every selected rule pack must exist and be published",
            status_code=409,
        )
    project_snapshot = await _project_snapshot(session, project.id)
    rule_snapshot = await _pack_snapshot(session, packs)
    input_hash = _canonical_hash({"project": project_snapshot, "rule_packs": rule_snapshot})
    package = ReviewPackage(
        project_id=project.id,
        name=payload.name,
        status="frozen",
        project_snapshot_hash=_canonical_hash(project_snapshot),
    )
    session.add(package)
    await session.flush()
    for pack in packs:
        session.add(ReviewPackageRulePack(review_package_id=package.id, rule_pack_id=pack.id))
    run = CheckRun(
        review_package_id=package.id,
        status="queued",
        project_snapshot=project_snapshot,
        rule_pack_snapshot=rule_snapshot,
        engine_version=ENGINE_VERSION,
        input_hash=input_hash,
    )
    session.add(run)
    await session.flush()
    job = Job(
        organization_id=actor.organization_id,
        project_id=project.id,
        job_type="check.run",
        status=JobStatus.QUEUED,
        progress=0,
        input_data={"check_run_id": str(run.id)},
        request_id=get_request_id(request),
    )
    session.add(job)
    record_audit_event(
        session,
        organization_id=actor.organization_id,
        actor_id=actor.user_id,
        action="check_run.queued",
        entity_type="check_run",
        entity_id=run.id,
        request_id=get_request_id(request),
        payload={"input_hash": input_hash, "rule_pack_ids": [str(item) for item in pack_ids]},
    )
    await session.commit()
    await dispatch_persisted_job(session, dispatcher, job)
    return CheckRunCreated(
        run=await _run_response(session, run), job=JobResponse.model_validate(job)
    )


@router.get("/check-runs/{run_id}", response_model=CheckRunResponse)
async def get_check_run(
    run_id: UUID,
    actor: Annotated[Actor, Depends(get_current_actor)],
    session: Annotated[AsyncSession, Depends(get_database_session)],
) -> CheckRunResponse:
    return await _run_response(session, await _owned_run(session, run_id, actor))


@router.get("/projects/{project_id}/check-runs", response_model=list[CheckRunResponse])
async def list_check_runs(
    project_id: UUID,
    actor: Annotated[Actor, Depends(get_current_actor)],
    session: Annotated[AsyncSession, Depends(get_database_session)],
) -> list[CheckRunResponse]:
    owned = await session.scalar(
        select(Project.id).where(
            Project.id == project_id, Project.organization_id == actor.organization_id
        )
    )
    if owned is None:
        raise ApplicationError("project_not_found", "Project was not found", status_code=404)
    runs = list(
        await session.scalars(
            select(CheckRun)
            .join(ReviewPackage)
            .where(ReviewPackage.project_id == project_id)
            .order_by(CheckRun.created_at.desc())
        )
    )
    return [await _run_response(session, item) for item in runs]


@router.get("/check-runs/{run_id}/export")
async def export_check_run(
    run_id: UUID,
    actor: Annotated[Actor, Depends(get_current_actor)],
    session: Annotated[AsyncSession, Depends(get_database_session)],
) -> JSONResponse:
    response = await _run_response(session, await _owned_run(session, run_id, actor))
    return JSONResponse(
        content=response.model_dump(mode="json"),
        headers={"Content-Disposition": f'attachment; filename="check-run-{run_id}.json"'},
    )


@router.get("/check-runs/{run_id}/workbench", response_model=WorkbenchResponse)
async def get_check_workbench(
    run_id: UUID,
    actor: Annotated[Actor, Depends(get_current_actor)],
    session: Annotated[AsyncSession, Depends(get_database_session)],
    storage: Annotated[ObjectStorage, Depends(get_object_storage)],
) -> WorkbenchResponse:
    return await _workbench_response(session, storage, await _owned_run(session, run_id, actor))


@router.patch("/check-results/{result_id}", response_model=CheckResultResponse)
async def update_check_result(
    result_id: UUID,
    payload: FindingUpdate,
    request: Request,
    actor: Annotated[Actor, Depends(get_current_actor)],
    session: Annotated[AsyncSession, Depends(get_database_session)],
) -> CheckResult:
    result = await session.scalar(
        select(CheckResult)
        .join(CheckRun)
        .join(ReviewPackage)
        .join(Project)
        .where(CheckResult.id == result_id, Project.organization_id == actor.organization_id)
    )
    if result is None:
        raise ApplicationError(
            "check_result_not_found", "Check result was not found", status_code=404
        )
    changes = payload.model_dump(exclude_unset=True)
    assignee_id = changes.get("assignee_id")
    if assignee_id is not None:
        assignee = await session.scalar(
            select(User).where(
                User.id == assignee_id, User.organization_id == actor.organization_id
            )
        )
        if assignee is None:
            raise ApplicationError("assignee_not_found", "Assignee was not found", status_code=404)
    for field, value in changes.items():
        setattr(result, field, value)
    record_audit_event(
        session,
        organization_id=actor.organization_id,
        actor_id=actor.user_id,
        action="check_result.updated",
        entity_type="check_result",
        entity_id=result.id,
        request_id=get_request_id(request),
        payload={
            key: str(value) if isinstance(value, UUID) else value
            for key, value in changes.items()
        },
    )
    await session.commit()
    await session.refresh(result)
    return result


def _xlsx_report(workbench: WorkbenchResponse) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Findings"
    sheet.append(
        [
            "Result ID",
            "Status",
            "Severity",
            "Workflow",
            "Message",
            "Drawing evidence",
            "Regulation evidence",
            "Reviewer notes",
        ]
    )
    for finding in workbench.findings:
        project_refs = "; ".join(
            f"page {item.location.get('page', '?')}: {item.excerpt or ''}"
            for item in finding.project_evidence
        )
        regulation_refs = "; ".join(
            f"page {item.location.get('page', '?')}: {item.excerpt or ''}"
            for item in finding.regulation_evidence
        )
        sheet.append(
            [
                str(finding.result_id),
                finding.status,
                finding.severity,
                finding.workflow_status,
                finding.message,
                project_refs,
                regulation_refs,
                finding.reviewer_notes or "",
            ]
        )
    sheet.freeze_panes = "A2"
    for index, column in enumerate(sheet.columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = min(
            60, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
        )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_report(workbench: WorkbenchResponse) -> bytes:
    buffer = io.BytesIO()
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    canvas = Canvas(buffer, pagesize=A4)
    _, height = A4
    canvas.setTitle(f"Compliance check {workbench.run_id}")
    canvas.setFont("STSong-Light", 15)
    canvas.drawString(42, height - 46, "Evidence-backed compliance review")
    y = height - 76
    for index, finding in enumerate(workbench.findings, start=1):
        lines = [
            f"{index}. [{finding.severity}] {finding.status} / {finding.workflow_status}",
            finding.message,
            "Drawing: "
            + "; ".join(
                f"p.{item.location.get('page', '?')} {item.excerpt or ''}"
                for item in finding.project_evidence
            ),
            "Regulation: "
            + "; ".join(
                f"p.{item.location.get('page', '?')} {item.excerpt or ''}"
                for item in finding.regulation_evidence
            ),
        ]
        for line in lines:
            if y < 54:
                canvas.showPage()
                y = height - 46
            canvas.setFont("STSong-Light", 10)
            canvas.drawString(42, y, line[:105])
            y -= 16
        y -= 7
    canvas.save()
    return buffer.getvalue()


@router.get("/check-runs/{run_id}/reports/{report_format}")
async def export_check_report(
    run_id: UUID,
    report_format: str,
    actor: Annotated[Actor, Depends(get_current_actor)],
    session: Annotated[AsyncSession, Depends(get_database_session)],
    storage: Annotated[ObjectStorage, Depends(get_object_storage)],
) -> StreamingResponse:
    if report_format not in {"pdf", "xlsx"}:
        raise ApplicationError(
            "unsupported_report_format", "Report format must be pdf or xlsx", status_code=404
        )
    workbench = await _workbench_response(
        session, storage, await _owned_run(session, run_id, actor)
    )
    content = _pdf_report(workbench) if report_format == "pdf" else _xlsx_report(workbench)
    media_type = (
        "application/pdf"
        if report_format == "pdf"
        else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return StreamingResponse(
        io.BytesIO(content),
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="check-run-{run_id}.{report_format}"'
        },
    )
