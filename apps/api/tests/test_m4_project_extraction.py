import asyncio
import io
import json
from pathlib import Path
from typing import Any
from uuid import UUID

import ifcopenshell
import pytest
from docx import Document
from openpyxl import Workbook

from app.api.v1.checks import _project_snapshot
from app.services.job_runtime import run_persisted_job
from app.services.pdf_extraction import PdfiumRegulationExtractor
from app.services.project_extraction import (
    FACT_DEFINITIONS,
    ExtractedSource,
    ProjectDocumentExtractor,
    classify_document,
    extract_pattern_facts,
)
from app.services.regulation_parser import ExtractedPage, TextLine
from app.tasks import project_extraction


def _xlsx_bytes(value: int) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fire Design"
    sheet.append(["Egress door clear width", value, "cm"])
    sheet.append(["Automatic sprinkler", "provided"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_fifteen_fact_definitions_extract_typed_values_and_locations() -> None:
    fixture_path = Path(__file__).parent / "fixtures" / "m4_golden_facts.json"
    golden = json.loads(fixture_path.read_text(encoding="utf-8"))
    text = "\n".join(golden["source"])
    facts = extract_pattern_facts([ExtractedSource(text, {"page": 3, "bbox": {"x0": 1}}, 0.93)])
    by_key = {fact.key: fact for fact in facts}
    assert len(FACT_DEFINITIONS) == 15
    assert set(by_key) == {item.key for item in FACT_DEFINITIONS}
    assert {key: fact.value for key, fact in by_key.items()} == golden["expected"]
    assert by_key["building.height_m"].value == 23.5
    assert by_key["building.floor_count"].value == 6
    assert by_key["egress.door_clear_width_m"].unit == "cm"
    assert by_key["smoke_control.provided"].value is False
    assert by_key["fire_alarm.provided"].location["page"] == 3


def test_docx_xlsx_and_ifc_extractors_use_precise_source_locations(
    tmp_path: Path, m1_environment: Any
) -> None:
    _client, storage, _dispatcher, _session_factory = m1_environment

    document = Document()
    document.add_paragraph("Building height: 18m")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Number of safety exits"
    table.cell(0, 1).text = "3"
    docx_stream = io.BytesIO()
    document.save(docx_stream)
    storage.objects["sample.docx"] = docx_stream.getvalue()

    storage.objects["sample.xlsx"] = _xlsx_bytes(120)

    model = ifcopenshell.file(schema="IFC4")
    for ifc_type, count in (
        ("IfcBuildingStorey", 2),
        ("IfcSpace", 3),
        ("IfcDoor", 4),
        ("IfcStair", 1),
    ):
        for _index in range(count):
            model.create_entity(ifc_type, GlobalId=ifcopenshell.guid.new())
    ifc_path = tmp_path / "sample.ifc"
    model.write(str(ifc_path))
    storage.objects["sample.ifc"] = ifc_path.read_bytes()

    extractor = ProjectDocumentExtractor(storage)
    docx_kind, docx_facts = extractor.extract("sample.docx", "brief.docx")
    xlsx_kind, xlsx_facts = extractor.extract("sample.xlsx", "schedule.xlsx")
    ifc_kind, ifc_facts = extractor.extract("sample.ifc", "model.ifc")

    assert docx_kind == "docx"
    assert {fact.key for fact in docx_facts} == {"building.height_m", "exit.count"}
    assert any(fact.location.get("paragraph") == 1 for fact in docx_facts)
    assert xlsx_kind == "xlsx"
    assert {fact.key for fact in xlsx_facts} == {
        "egress.door_clear_width_m",
        "sprinkler.provided",
    }
    assert xlsx_facts[0].location["sheet"] == "Fire Design"
    assert ifc_kind == "ifc"
    assert {fact.key: fact.value for fact in ifc_facts[:4]} == {
        "building.floor_count": 2,
        "space.count": 3,
        "door.count": 4,
        "stair.count": 1,
    }
    assert classify_document("MODEL.IFC") == "ifc"


def test_pdf_project_extraction_reuses_positioned_m2_output(
    m1_environment: Any, monkeypatch: Any
) -> None:
    _client, storage, _dispatcher, _session_factory = m1_environment
    pages = [
        ExtractedPage(
            page_number=4,
            width=595,
            height=842,
            method="text_layer",
            lines=[
                TextLine(
                    "Building height: 21m",
                    {"x0": 10, "y0": 20, "x1": 120, "y1": 32},
                    0.98,
                )
            ],
            image_png=b"png",
        )
    ]
    monkeypatch.setattr(PdfiumRegulationExtractor, "extract", lambda _self, _key: pages)

    kind, facts = ProjectDocumentExtractor(storage).extract("sample.pdf", "brief.pdf")

    assert kind == "pdf"
    assert facts[0].key == "building.height_m"
    assert facts[0].location == {
        "page": 4,
        "bbox": {"x0": 10, "y0": 20, "x1": 120, "y1": 32},
        "method": "text_layer",
    }
    assert facts[0].confidence == 0.98
    with pytest.raises(ValueError, match="Unsupported project document"):
        classify_document("drawing.dwg")


def test_extraction_candidates_conflict_and_require_explicit_decisions(
    m1_environment: Any, monkeypatch: Any
) -> None:
    client, storage, dispatcher, session_factory = m1_environment
    project = client.post("/api/v1/projects", json={"name": "M4 pilot"}).json()
    monkeypatch.setattr(project_extraction, "get_object_storage", lambda: storage)

    def upload_and_extract(value: int) -> str:
        upload = client.post(
            f"/api/v1/projects/{project['id']}/files",
            data={"logical_name": f"Fire schedule {value}"},
            files={
                "upload": (
                    f"schedule-{value}.xlsx",
                    _xlsx_bytes(value),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert upload.status_code == 201
        file_version_id = upload.json()["file_version"]["id"]
        queued = client.post(
            f"/api/v1/projects/{project['id']}/extractions",
            json={"file_version_id": file_version_id},
        )
        assert queued.status_code == 201
        assert queued.json()["document_kind"] == "xlsx"
        job_id, job_type = dispatcher.dispatched[-1]
        assert job_type == "project.extract"
        asyncio.run(
            run_persisted_job(
                session_factory,
                job_id,
                project_extraction.extract_project_facts,
                error_code="project_extraction_failed",
            )
        )
        return file_version_id

    first_file_id = upload_and_extract(90)
    second_file_id = upload_and_extract(120)
    assert first_file_id != second_file_id
    non_pdf_regulation = client.post(
        "/api/v1/regulations/ingestions",
        json={
            "file_version_id": first_file_id,
            "code": "TEST-XLSX",
            "title": "Invalid spreadsheet regulation",
            "edition": "2026",
            "jurisdiction": "Test",
        },
    )
    assert non_pdf_regulation.status_code == 415
    assert non_pdf_regulation.json()["error"]["code"] == "regulation_pdf_required"

    candidates = client.get(f"/api/v1/projects/{project['id']}/fact-candidates").json()
    widths = [item for item in candidates if item["key"] == "egress.door_clear_width_m"]
    assert {item["verification_status"] for item in widths} == {"conflicting"}
    assert {item["value"] for item in widths} == {90.0, 120.0}
    assert all(item["evidence"][0]["file_version_id"] for item in widths)
    assert all(item["evidence"][0]["location"]["sheet"] == "Fire Design" for item in widths)

    selected = next(item for item in widths if item["value"] == 120.0)
    rejected = next(item for item in widths if item["value"] == 90.0)
    verified_response = client.post(
        f"/api/v1/fact-candidates/{selected['id']}/verify",
        json={"reason": "Architect checked the latest schedule"},
    )
    assert verified_response.status_code == 200
    assert verified_response.json()["verification_status"] == "verified"
    reject_response = client.post(
        f"/api/v1/fact-candidates/{rejected['id']}/reject",
        json={"reason": "Superseded schedule"},
    )
    assert reject_response.status_code == 200
    assert reject_response.json()["verification_status"] == "rejected"
    assert (
        client.post(
            f"/api/v1/fact-candidates/{rejected['id']}/reject",
            json={"reason": "Second decision"},
        ).status_code
        == 409
    )

    verified_facts = client.get(f"/api/v1/projects/{project['id']}/facts").json()
    verified_width = next(
        item for item in verified_facts if item["key"] == "egress.door_clear_width_m"
    )
    assert verified_width["value"] == 120.0
    assert verified_width["source"] == "spreadsheet"
    assert len(client.get("/api/v1/fact-types").json()) == 15

    async def load_m3_snapshot() -> dict[str, Any]:
        async with session_factory() as session:
            return await _project_snapshot(session, UUID(project["id"]))

    snapshot = asyncio.run(load_m3_snapshot())
    snapshot_width = next(
        item for item in snapshot["facts"] if item["key"] == "egress.door_clear_width_m"
    )
    assert snapshot_width["value"] == 120.0
    assert snapshot_width["evidence_ids"]
